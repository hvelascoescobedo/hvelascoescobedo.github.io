#!/usr/bin/env python3
"""Revisor rapido de reportes: avisa cuales hay que volver a descargar.

Programa aparte del buscador de items.  Su unico trabajo es abrir cada
reporte de una carpeta y decir si trae productos adentro o si llego vacio
(la descarga salio mal), sin hacer ningun analisis de ventas.

En Windows basta con dar doble clic en "Revisar archivos.bat".  Tambien:

    python revisar_archivos.py
    python revisar_archivos.py --carpeta "C:\\Users\\...\\Hora Bostons"

Es rapido a proposito: en vez de leer el reporte completo como hace el
buscador (unos 5 segundos por PDF), lee pagina por pagina y se detiene en
cuanto encuentra el primer producto, que casi siempre esta en la pagina 1
(unas decimas de segundo por archivo).
"""

from __future__ import annotations

import argparse
import logging
import re
import sys
import time
from pathlib import Path

# pypdf avisa cosas como "EOF marker not found" al abrir un PDF dañado; aqui
# eso ya se reporta con un mensaje propio, asi que se calla el ruido interno.
logging.getLogger("pypdf").setLevel(logging.CRITICAL)

CARPETA_POR_DEFECTO = "datos"
EXTENSIONES_VALIDAS = {".txt", ".pdf", ".prn", ".rpt", ".dat", ".text"}
# Descargas a medio bajar, segun el navegador; no son reportes todavia.
EXTENSIONES_A_MEDIAS = {".crdownload", ".part", ".partial", ".tmp", ".download", ".opdownload"}
COMILLAS = "\"'\u201c\u201d\u2018\u2019\u00ab\u00bb"

# Renglon de producto: sangria, numero de item y luego el nombre.
RE_ITEM = re.compile(r"^\s{2,}\d+\s", re.MULTILINE)
# "Category Totals": si aparece, el reporte si llego hasta el final.
RE_TOTALES = re.compile(r"^Category Totals", re.MULTILINE)
# "For the period:  8/10/26 to 8/13/26"
RE_PERIODO = re.compile(
    r"For the period:\s*(\d{1,2}/\d{1,2}/\d{2,4})\s*to\s*(\d{1,2}/\d{1,2}/\d{2,4})"
)
# Los primeros 4 digitos del nombre del archivo son el restaurante.
RE_ID = re.compile(r"^\s*(\d{4})(?!\d)")

ESTADO_OK = "OK"
ESTADO_VACIO = "VACIO"
ESTADO_NO_ABRE = "NO ABRE"


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------


def limpiar_ruta(texto: str) -> str:
    """Acepta la ruta con o sin comillas, como la copia Windows."""
    limpio = texto.strip()
    while len(limpio) >= 2 and limpio[0] in COMILLAS and limpio[-1] in COMILLAS:
        limpio = limpio[1:-1].strip()
    return limpio.strip(COMILLAS).strip()


def restaurante_de(ruta: Path) -> str:
    """#ID de restaurante segun el nombre del archivo ('6004_10_al_13' -> 6004)."""
    coincidencia = RE_ID.match(ruta.stem)
    return coincidencia.group(1) if coincidencia else ""


def _fecha(texto: str) -> str:
    """'8/10/26' (mes/dia/anio del reporte) -> '10/08/2026'."""
    partes = texto.split("/")
    if len(partes) != 3:
        return texto
    mes, dia, anio = partes
    anio = anio if len(anio) == 4 else f"20{anio}"
    return f"{int(dia):02d}/{int(mes):02d}/{anio}"


def periodo_de(texto: str) -> str:
    """Periodo que viene dentro del reporte, en formato dia/mes/anio."""
    coincidencia = RE_PERIODO.search(texto)
    if not coincidencia:
        return ""
    return f"del {_fecha(coincidencia.group(1))} al {_fecha(coincidencia.group(2))}"


def cargar_diccionario(lugares) -> dict[str, str]:
    """Catalogo #ID -> nombre del restaurante; opcional, si se puede leer."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {}

    for carpeta in lugares:
        candidatos = sorted(Path(carpeta).glob("*iccionario*.xlsx")) if carpeta.is_dir() else []
        for candidato in candidatos:
            if candidato.name.startswith("~$"):
                continue
            try:
                libro = load_workbook(candidato, data_only=True, read_only=True)
            except Exception:
                continue
            mapa: dict[str, str] = {}
            try:
                for hoja in libro.worksheets:
                    for fila in hoja.iter_rows(values_only=True):
                        for celda in fila:
                            texto = str(celda).strip() if celda is not None else ""
                            encontrado = RE_ID.match(texto)
                            if encontrado:
                                nombre = texto[encontrado.end():].strip(" -\u2013\u2014:\t")
                                mapa.setdefault(encontrado.group(1), nombre)
                                break
            finally:
                libro.close()
            if mapa:
                return mapa
    return {}


# ---------------------------------------------------------------------------
# Revision de un archivo
# ---------------------------------------------------------------------------


def _texto_de_txt(ruta: Path) -> str:
    datos = ruta.read_bytes()
    for codificacion in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return datos.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return datos.decode("latin-1", errors="replace")


def _revisar_pdf(ruta: Path) -> tuple[str, str, str]:
    """Revisa un PDF pagina por pagina, parando en cuanto haya un producto."""
    from pypdf import PdfReader

    lector = PdfReader(str(ruta))
    if lector.is_encrypted:
        try:
            lector.decrypt("")
        except Exception:
            return ESTADO_NO_ABRE, "el PDF esta protegido con contrasena", ""

    periodo = ""
    con_totales = False
    for pagina in lector.pages:
        # Texto plano (no "layout"): aqui solo interesa SI hay productos, no
        # en que columna cae cada dato, y asi es varias veces mas rapido.
        texto = pagina.extract_text() or ""
        if not periodo:
            periodo = periodo_de(texto)
        if RE_ITEM.search(texto):
            return ESTADO_OK, "", periodo  # ya hay datos: no se lee el resto
        con_totales = con_totales or bool(RE_TOTALES.search(texto))

    if con_totales:
        return ESTADO_VACIO, "el reporte llego sin un solo producto", periodo
    return ESTADO_VACIO, "no se encontro ningun producto ni el corte de totales", periodo


def _revisar_txt(ruta: Path) -> tuple[str, str, str]:
    texto = _texto_de_txt(ruta)
    periodo = periodo_de(texto)
    if RE_ITEM.search(texto):
        return ESTADO_OK, "", periodo
    if RE_TOTALES.search(texto):
        return ESTADO_VACIO, "el reporte llego sin un solo producto", periodo
    return ESTADO_VACIO, "no se encontro ningun producto ni el corte de totales", periodo


def revisar(ruta: Path) -> dict:
    """Revisa un archivo y devuelve su estado, motivo y periodo."""
    registro = {
        "ruta": ruta,
        "restaurante": restaurante_de(ruta),
        "estado": ESTADO_NO_ABRE,
        "motivo": "",
        "periodo": "",
    }
    try:
        if ruta.stat().st_size == 0:
            # Ojo: "Microsoft Print to PDF" crea el archivo en 0 bytes y lo
            # llena despues, asi que esto puede ser solo que aun no termina.
            registro["motivo"] = (
                "pesa 0 bytes; si lo acabas de mandar a 'Microsoft Print to PDF', "
                "espera a que Windows termine de guardarlo y revisalo otra vez"
            )
            registro["estado"] = ESTADO_VACIO
            return registro

        if ruta.suffix.lower() == ".pdf":
            estado, motivo, periodo = _revisar_pdf(ruta)
        else:
            estado, motivo, periodo = _revisar_txt(ruta)
        registro.update(estado=estado, motivo=motivo, periodo=periodo)
    except ImportError:
        registro["motivo"] = "falta la libreria 'pypdf' para leer PDFs"
    except Exception as error:
        registro["motivo"] = f"{type(error).__name__}: {error}"
    return registro


# ---------------------------------------------------------------------------
# Programa
# ---------------------------------------------------------------------------


def pedir_carpeta(inicial: str | None, base: Path) -> Path:
    candidata = limpiar_ruta(inicial) if inicial else ""
    while True:
        if not candidata:
            try:
                candidata = limpiar_ruta(
                    input(f"Carpeta con los reportes [{CARPETA_POR_DEFECTO}]: ")
                ) or CARPETA_POR_DEFECTO
            except EOFError:
                candidata = CARPETA_POR_DEFECTO
        carpeta = Path(candidata).expanduser()
        if not carpeta.is_absolute() and not carpeta.exists():
            carpeta = base / candidata
        if carpeta.is_dir():
            return carpeta
        print(f"  No encontre la carpeta '{candidata}'.")
        candidata = ""


def listar_archivos(carpeta: Path, recursivo: bool = True) -> list[Path]:
    patron = "**/*" if recursivo else "*"
    archivos = [
        p
        for p in carpeta.glob(patron)
        if p.is_file()
        and p.suffix.lower() in EXTENSIONES_VALIDAS
        and not p.name.startswith((".", "~$"))
    ]
    return sorted(archivos, key=lambda p: (p.parent.as_posix(), p.name))


def descargas_en_curso(carpeta: Path, recursivo: bool = True) -> list[Path]:
    """Archivos a medio bajar (.crdownload de Chrome, .part de Firefox, etc.)."""
    patron = "**/*" if recursivo else "*"
    return sorted(
        p for p in carpeta.glob(patron) if p.is_file() and p.suffix.lower() in EXTENSIONES_A_MEDIAS
    )


def firma(ruta: Path) -> tuple[int, int]:
    """Tamano y fecha del archivo; si cambia, es que lo volvieron a descargar."""
    try:
        estado = ruta.stat()
        return (estado.st_size, estado.st_mtime_ns)
    except OSError:
        return (-1, -1)


def esperar_a_que_termine(
    ruta: Path, pausa: float = 0.4, intentos: int = 12, intentos_en_cero: int = 30
) -> None:
    """Espera a que el archivo deje de crecer, por si sigue escribiendose.

    Sin esto, dar Enter demasiado pronto haria leer un archivo a medias y
    reportar un problema que en realidad no existe.

    Un archivo en 0 bytes NO cuenta como terminado, y por eso se le da mucho
    mas tiempo: "Microsoft Print to PDF" crea el archivo vacio en cuanto se
    elige donde guardarlo, y el spooler lo llena varios segundos despues.
    """
    anterior = firma(ruta)
    limite = intentos_en_cero if anterior[0] == 0 else intentos
    for _ in range(limite):
        time.sleep(pausa)
        actual = firma(ruta)
        if actual == anterior and actual[0] > 0:
            return  # dejo de crecer y ya trae contenido
        anterior = actual


def _imprimir_resultado(registro: dict, prefijo: str = "  ") -> None:
    """Un renglon con el estado del archivo y, si aplica, el motivo."""
    print(f"{prefijo}{registro['estado']:<8} {registro['ruta'].name}")
    if registro["motivo"]:
        print(f"{' ' * len(prefijo)}         -> {registro['motivo']}")


def _imprimir_pendientes(problemas: list[dict], diccionario: dict[str, str]) -> None:
    """Lista de los reportes que hay que volver a descargar."""
    print("\nHay que volver a descargar estos reportes:\n")
    for registro in sorted(problemas, key=lambda r: (r["restaurante"] or "~", r["ruta"].name)):
        etiqueta = registro["restaurante"] or "(sin ID)"
        nombre = diccionario.get(registro["restaurante"], "")
        print(
            f"  {f'{etiqueta} {nombre}'.strip():<32} {registro['periodo'] or '(sin fechas)':<26} "
            f"{registro['ruta'].name}"
        )
    restaurantes = sorted({r["restaurante"] for r in problemas if r["restaurante"]})
    if restaurantes:
        print(f"\nRestaurantes a re-descargar: {', '.join(restaurantes)}")


def vigilar(
    carpeta: Path,
    diccionario: dict[str, str],
    recursivo: bool = True,
    leer_orden=input,
    pausa: float = 0.4,
) -> int:
    """Revisa lo que ya hay y luego espera, revisando cada archivo nuevo.

    Pensado para ir descargando los reportes de uno en uno: se descarga uno,
    se presiona Enter y el programa dice al momento si ese archivo esta bien,
    sin volver a revisar los que ya reviso.
    """
    conocidos: dict[Path, tuple[int, int]] = {}   # archivo -> firma ya revisada
    ultimo: dict[Path, dict] = {}                 # archivo -> como quedo

    def revisar_lote(rutas: list[Path], titulo: str) -> None:
        if not rutas:
            return
        print(f"\n{titulo}\n")
        for ruta in rutas:
            esperar_a_que_termine(ruta, pausa=pausa)
            registro = revisar(ruta)
            conocidos[ruta] = firma(ruta)
            ultimo[ruta] = registro
            _imprimir_resultado(registro, prefijo="  ")
        malos = [r for r in (ultimo[ruta] for ruta in rutas) if r["estado"] != ESTADO_OK]
        if malos:
            print(f"\n  >>> OJO: {len(malos)} de {len(rutas)} hay que volver a descargarlo.")
        elif len(rutas) > 1:
            print(f"\n  Los {len(rutas)} archivos estan bien.")
        else:
            print("\n  Ese archivo esta bien.")

    revisar_lote(listar_archivos(carpeta, recursivo), "Primero reviso lo que ya hay en la carpeta:")
    if not conocidos:
        print("\nLa carpeta esta vacia por ahora; en cuanto descargues algo, dale Enter.")

    print("\n" + "-" * 70)
    print("Ya puedes descargar los reportes de uno en uno.")
    print("  Enter  -> reviso lo que acabe de llegar")
    print("  todo   -> vuelvo a revisar todos los archivos de la carpeta")
    print("  salir  -> termino y te doy el resumen")
    print("-" * 70)

    while True:
        try:
            orden = (leer_orden("\nEsperando... (Enter para revisar) > ") or "").strip().lower()
        except (EOFError, KeyboardInterrupt):
            print()
            break

        if orden in {"salir", "s", "q", "quit", "exit", "no", "n"}:
            break

        pendientes_descarga = descargas_en_curso(carpeta, recursivo)
        if pendientes_descarga:
            print(f"\n  Espera: hay {len(pendientes_descarga)} descarga(s) todavia en curso:")
            for ruta in pendientes_descarga:
                print(f"    {ruta.name}")
            print("  Cuando terminen, dale Enter otra vez.")
            continue

        if orden in {"todo", "t", "todos"}:
            conocidos.clear()
            revisar_lote(listar_archivos(carpeta, recursivo), "Reviso otra vez todos los archivos:")
            continue

        # Nuevos: los que no se habian visto, o los que cambiaron porque se
        # volvieron a descargar encima del anterior.
        nuevos = [
            ruta
            for ruta in listar_archivos(carpeta, recursivo)
            if conocidos.get(ruta) != firma(ruta)
        ]
        if not nuevos:
            print("\n  No llego ningun archivo nuevo todavia.")
            continue

        repetidos = {ruta for ruta in nuevos if ruta in conocidos}
        titulo = f"Reviso {len(nuevos)} archivo(s) nuevo(s):"
        if repetidos:
            titulo = f"Reviso {len(nuevos)} archivo(s) ({len(repetidos)} re-descargado/s):"
        revisar_lote(nuevos, titulo)

    # Resumen final de todo lo revisado en la sesion.
    revisados = list(ultimo.values())
    problemas = [r for r in revisados if r["estado"] != ESTADO_OK]
    print()
    print("=" * 70)
    if not revisados:
        print("No se reviso ningun archivo.")
        print("=" * 70)
        return 1
    if not problemas:
        cuantos = (
            "el archivo revisado trae datos"
            if len(revisados) == 1
            else f"los {len(revisados)} archivos revisados traen datos"
        )
        print(f"TODO BIEN: {cuantos}.")
        print("=" * 70)
        print("\nYa puedes correr el buscador de items.")
        return 0

    print(f"RESULTADO: {len(revisados) - len(problemas)} bien, {len(problemas)} con problema")
    print("=" * 70)
    _imprimir_pendientes(problemas, diccionario)
    return 1


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Revisa que los reportes de una carpeta no esten vacios."
    )
    parser.add_argument("--carpeta", help="Carpeta con los reportes (por defecto: datos)")
    parser.add_argument(
        "--no-recursivo", action="store_true", help="No entrar a las subcarpetas"
    )
    parser.add_argument(
        "--solo-problemas", action="store_true", help="Listar unicamente lo que salio mal"
    )
    parser.add_argument(
        "--vigilar",
        action="store_true",
        help="No cerrar: revisar cada archivo nuevo conforme se vaya descargando",
    )
    argumentos = parser.parse_args(argv)

    base = Path(__file__).resolve().parent
    if getattr(sys, "frozen", False):  # ejecutable hecho con PyInstaller
        base = Path(sys.executable).resolve().parent

    print()
    print("=" * 70)
    if argumentos.vigilar:
        print("REVISION CONTINUA  (reviso cada archivo conforme lo descargas)")
    else:
        print("REVISION DE ARCHIVOS  (¿estan completos los reportes?)")
    print("=" * 70)

    carpeta = pedir_carpeta(argumentos.carpeta, base)
    print(f"Carpeta: {carpeta}")

    lugares = [carpeta, base, base / "sistema"]
    if getattr(sys, "frozen", False):  # copia incluida dentro del .exe
        lugares.append(Path(getattr(sys, "_MEIPASS", base)))
    diccionario = cargar_diccionario(lugares)

    if argumentos.vigilar:
        return vigilar(carpeta, diccionario, recursivo=not argumentos.no_recursivo)

    archivos = listar_archivos(carpeta, recursivo=not argumentos.no_recursivo)
    if not archivos:
        print("\nNo hay archivos de reporte (.txt o .pdf) en esa carpeta.")
        return 1

    print(f"\nRevisando {len(archivos)} archivo(s)...\n")
    revisados = []
    for numero, ruta in enumerate(archivos, start=1):
        registro = revisar(ruta)
        revisados.append(registro)
        if argumentos.solo_problemas and registro["estado"] == ESTADO_OK:
            continue
        print(f"  [{numero:>3}/{len(archivos)}]  {registro['estado']:<8} {ruta.name}")
        if registro["motivo"]:
            print(f"                       -> {registro['motivo']}")

    problemas = [r for r in revisados if r["estado"] != ESTADO_OK]
    buenos = len(revisados) - len(problemas)

    print()
    print("=" * 70)
    if not problemas:
        print(f"TODO BIEN: los {buenos} archivos traen datos.")
        print("=" * 70)
        print("\nYa puedes correr el buscador de items.")
        return 0

    print(f"RESULTADO: {buenos} bien, {len(problemas)} con problema")
    print("=" * 70)
    _imprimir_pendientes(problemas, diccionario)
    return 1


if __name__ == "__main__":
    codigo = main()
    if "--vigilar" not in sys.argv and sys.stdin and sys.stdin.isatty():
        try:
            input("\nPresiona Enter para salir...")
        except EOFError:
            pass
    raise SystemExit(codigo)
